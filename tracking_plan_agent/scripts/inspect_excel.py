# -*- coding: utf-8 -*-
"""检查 Excel 结构并导出内容，供埋点 agent 解析需求文档与全量埋点表使用。

用法：
    python inspect_excel.py <文件.xlsx> [--sheet Sheet1] [--max-rows 200] [--json]

输出：sheet 列表、表头行、逐行内容（Markdown 表格，默认；--json 输出 JSON）。
说明：
    - 以 data_only=True 打开，取公式计算值；DISPIMG 公式会原样显示为公式文本。
    - 表头行自动探测：前 10 行中含有"指标名"或"埋点"或"需求"等关键词的行视为表头。
"""
import argparse
import json
import sys

import openpyxl

HEADER_HINTS = ["指标名", "埋点", "需求", "事件", "指标", "设计", "端", "触发"]


def find_header_row(rows, max_scan=10):
    """返回表头所在行号（1-based）。启发式：命中提示词最多的行。"""
    best_row, best_score = 1, 0
    for i, row in enumerate(rows[:max_scan], start=1):
        score = 0
        for cell in row:
            if cell is None:
                continue
            text = str(cell)
            if any(h in text for h in HEADER_HINTS):
                score += 1
        if score > best_score:
            best_row, best_score = i, score
    return best_row


def cell_text(v):
    if v is None:
        return ""
    return str(v)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("path", help="Excel 文件路径")
    ap.add_argument("--sheet", help="sheet 名称，默认第一个非空 sheet")
    ap.add_argument("--max-rows", type=int, default=200, help="最多输出行数")
    ap.add_argument("--json", action="store_true", help="以 JSON 输出")
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.path, data_only=True)

    # 选 sheet：优先用户指定，否则取非空 sheet
    if args.sheet:
        if args.sheet not in wb.sheetnames:
            print(f"ERROR: sheet '{args.sheet}' 不存在，可用: {wb.sheetnames}", file=sys.stderr)
            sys.exit(1)
        ws = wb[args.sheet]
    else:
        ws = None
        for name in wb.sheetnames:
            candidate = wb[name]
            if candidate.max_row > 1 or candidate.max_column > 1:
                ws = candidate
                break
        if ws is None:
            ws = wb.worksheets[0]

    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    # 去掉尾部全空行
    while rows and all(c is None or str(c).strip() == "" for c in rows[-1]):
        rows.pop()

    header_row = find_header_row(rows)
    headers = [cell_text(c) for c in rows[header_row - 1]] if len(rows) >= header_row else []
    data_rows = rows[header_row:] if len(rows) > header_row else []

    result = {
        "file": args.path,
        "sheet": ws.title,
        "all_sheets": wb.sheetnames,
        "header_row": header_row,
        "headers": headers,
        "merged_cells": [str(r) for r in ws.merged_cells.ranges][:50],
        "row_count": len(data_rows),
        "rows": [],
    }

    for idx, row in enumerate(data_rows[: args.max_rows], start=header_row + 1):
        cells = [cell_text(c) for c in row]
        if any(c.strip() for c in cells):
            result["rows"].append({"row": idx, "cells": cells})

    if args.json:
        print(json.dumps(result, ensure_ascii=False, indent=1))
    else:
        print(f"文件: {result['file']}")
        print(f"Sheet: {result['sheet']}（全部: {result['all_sheets']}）")
        print(f"表头行: 第 {header_row} 行")
        print(f"数据行数: {result['row_count']}（截断至 {args.max_rows}）")
        if result["merged_cells"]:
            print(f"合并单元格: {result['merged_cells'][:10]}")
        print()
        # Markdown 表格
        print("| 行号 | " + " | ".join(f"C{i+1}:{h}" for i, h in enumerate(headers)) + " |")
        print("|---" * (len(headers) + 1) + "|")
        for r in result["rows"]:
            print(f"| {r['row']} | " + " | ".join(c.replace("\n", "⏎")[:120] for c in r["cells"]) + " |")
        if result["row_count"] > args.max_rows:
            print(f"\n⚠️ 仅显示前 {args.max_rows} 行，共 {result['row_count']} 行，可用 --max-rows 调整")


if __name__ == "__main__":
    main()
